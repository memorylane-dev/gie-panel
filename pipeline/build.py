# -*- coding: utf-8 -*-
import sys, os, json, math
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import ROOT, RAW_W1, RAW_W2, MAPPING, DELIVERY, N, ls, find_file
from taxonomy import TAXO, CATEGORIES, COMPETENCY, REVERSE_ITEMS, SCALE_MISFIT, MULTIRESP_FILL0, DATA_ISSUE, COST_ITEMS, BASE_ALIGN, SUBSCALE, CONDITIONAL_BASE
import openpyxl, pyreadstat, pandas as pd, numpy as np

OUT = DELIVERY
os.makedirs(OUT, exist_ok=True)

WAVES = [("W1", 1, 2021, "1주기"), ("W2", 2, 2025, "2주기")]
d1, d2 = RAW_W1, RAW_W2
SRC = {
 ("학생",1): find_file(d1,"student total"), ("학생",2): find_file(d2,"1. 학생"),
 ("학부모",1): find_file(d1,"PAR"),          ("학부모",2): find_file(d2,"2. 학부모.sav"),
 ("교사",1): find_file(d1,"TEA"),            ("교사",2): find_file(d2,"3. 교사.sav"),
}
IDCOL = {"학생":"STUID","학부모":"STUID","교사":"TID"}

DATA, META = {}, {}
for k,f in SRC.items():
    df, m = pyreadstat.read_sav(f)
    m.column_names_to_labels = {c:N(v or "") for c,v in m.column_names_to_labels.items()}
    m.variable_value_labels   = {c:{k2:N(v2) for k2,v2 in d.items()} for c,d in m.variable_value_labels.items()}
    DATA[k], META[k] = df, m

# 학교 지역규모
db1,_ = pyreadstat.read_sav(find_file(d1,"SCH DB"))
db2,_ = pyreadstat.read_sav(find_file(d2,"5. SCH DB"))
REGION = {1: db1.set_index("SCHID")["YM1_DB0_3"].to_dict(),
          2: db2.set_index("SCHID")["YM2_DB0_3"].to_dict()}
REGION_LBL = {1.0:"중소도시", 2.0:"읍면지역"}

# ── 1. 지표 마스터 ────────────────────────────────────────────
wb = openpyxl.load_workbook(find_file(MAPPING,"변수매칭"), data_only=True)
raw = [[N(str(x)) if x is not None else "" for x in r] for r in wb["6_최종연계목록"].iter_rows(values_only=True)][2:]
raw = [r for r in raw if r[0]]

POS_KW = ["예","참여","해당","있음"]
def stat_type(nsel):
    return {"5→5":"likert","4→4":"likert","6→6":"likert","2→2":"binary","연속형":"continuous"}.get(nsel,"excluded")

def vlabels(resp, wave, var):
    if not var or var=="—": return {}
    return META[(resp,wave)].variable_value_labels.get(var, {})

ind_rows, vlab_rows, unmapped = [], [], []
for r in raw:
    cid,resp,gu,area,qname,sub,v1,v2,jud,nsel,cross,chk,note = r[:13]
    key = (resp, area, qname)
    if key not in TAXO:
        unmapped.append(key); continue
    cid_c, cname, sid, sname = TAXO[key]
    if cid in SUBSCALE: sid, sname = SUBSCALE[cid]
    st = stat_type(nsel)
    l1, l2 = vlabels(resp,1,v1), vlabels(resp,2,v2)
    codes = sorted(set(l1)|set(l2))
    pos_val = ""
    if st=="binary":
        for c in codes:
            lab = l2.get(c) or l1.get(c) or ""
            if any(k in lab for k in POS_KW): pos_val = c
    smin = min(codes) if codes else ""
    smax = max(codes) if codes else ""
    ts = "Y" if jud.startswith("①") else ("조건부" if jud.startswith("②") else "N")
    issue = DATA_ISSUE.get(cid,"")
    if issue: ts = "N"
    if cid in BASE_ALIGN:
        ts = "조건부"
        issue = BASE_ALIGN[cid][4]
    if cid in COST_ITEMS:
        ts = "조건부"
        issue = ("1주기는 미지출을 0으로 기입, 2주기는 공란 처리 — 전체 평균은 응답 기저가 다르다. "
                 "시계열 비교에는 가정이 필요 없는 '지출자평균/지출자중앙값'을 사용할 것. "
                 "'평균'·'지출참여율'은 2주기 결측=미지출 가정을 전제한다.")
    if st=="excluded": ts = "N"
    ind_rows.append(dict(
        indicator_id=cid, 대분류코드=cid_c, 대분류=cname, 소주제코드=sid, 소주제=sname,
        응답주체=resp, 원영역=f"{gu}>{area}", 문항명=qname, 하위문항=sub,
        지표명=f"{qname} - {sub}" if sub else qname,
        var_2021=v1, var_2025=v2, 연계판정=jud, 통계유형=st, 선지수=nsel,
        척도최소=smin, 척도최대=smax, 척도최소라벨=(l2.get(smin) or l1.get(smin) or ""),
        척도최대라벨=(l2.get(smax) or l1.get(smax) or ""),
        긍정코드=pos_val,
        역문항="Y" if cid in REVERSE_ITEMS else ("척도 부적합" if cid in SCALE_MISFIT else "N"),
        시계열비교가능=ts, 주체간비교=cross, 확인필요=chk,
        데이터이슈=issue, 응답기저=CONDITIONAL_BASE.get(cid,""),
        주의사항=note, 출처="변수매칭_6_최종연계목록"))
    for w,lab in ((2021,l1),(2025,l2)):
        for c,t in sorted(lab.items()):
            vlab_rows.append(dict(indicator_id=cid, 조사연도=w, 코드=c, 라벨=t))

# 학생 핵심역량 지표 추가
for cid, nm, v1, v2, jud, note in COMPETENCY:
    ts = "Y" if jud.startswith("①") else ("조건부" if jud.startswith("②") else "N")
    ind_rows.append(dict(
        indicator_id=cid, 대분류코드="C01", 대분류="학생성장", 소주제코드="S0101", 소주제="역량점수",
        응답주체="학생", 원영역="성과>학생 핵심역량", 문항명="학생 핵심역량 산출점수", 하위문항=nm,
        지표명=nm, var_2021=v1, var_2025=v2, 연계판정=jud, 통계유형="continuous", 선지수="연속형",
        척도최소="", 척도최대="", 척도최소라벨="", 척도최대라벨="", 긍정코드="", 역문항="N",
        시계열비교가능=ts, 주체간비교="", 확인필요="", 데이터이슈="", 응답기저="", 주의사항=note,
        출처="변수매칭_참고_미분류변수"))

dim = pd.DataFrame(ind_rows)
print("지표 수:", len(dim), "| 미매핑 문항:", set(unmapped))

# ── 2. 마이크로데이터(long) ───────────────────────────────────
recs = []
for _, ix in dim.iterrows():
    resp, cid = ix["응답주체"], ix["indicator_id"]
    for wcode, wnum, year, wlab in WAVES:
        var = ix["var_2021"] if wnum==1 else ix["var_2025"]
        if not var or var=="—": continue
        df, m = DATA[(resp,wnum)], META[(resp,wnum)]
        if var not in df.columns: continue
        s = df[var].copy()
        # R1: 다중응답 이분 — 미선택 결측을 0으로 보정
        fill_list = MULTIRESP_FILL0["wave1"] if wnum==1 else MULTIRESP_FILL0["wave2"]
        filled = False
        if ix["통계유형"]=="binary" and var in fill_list:
            s = s.fillna(0.0); filled = True
        # R2: 라벨 정의 범위 밖 값 결측 처리
        #  주의: R1로 보정한 0은 라벨에 정의돼 있지 않을 수 있으므로 허용 코드에 포함해야 한다.
        #  (그러지 않으면 R1이 무효화되어 1주기 다중응답 문항의 긍정응답률이 100%가 된다)
        lab = m.variable_value_labels.get(var, {})
        if lab and ix["통계유형"] in ("likert","binary"):
            allowed = set(lab.keys()) | ({0.0} if filled else set())
            s = s.where(s.isin(list(allowed)) | s.isna())
        # R7: 응답 기저 정렬 — 한 주기에만 분기가 걸린 문항의 기저를 맞춘다
        aligned = ""
        if cid in BASE_ALIGN:
            a_resp, a_wave, a_var, a_val, a_note = BASE_ALIGN[cid]
            if resp == a_resp and wnum == a_wave and a_var in df.columns:
                s = s.where(df[a_var] == a_val)
                aligned = f"기저정렬({a_var}={a_val:g})"
        # R3: 연속형 음수 결측
        if ix["통계유형"]=="continuous":
            s = s.where(s >= 0)
        sub = pd.DataFrame({
            "조사명": "경기학교교육실태조사", "학교급": "중학교",
            "조사연도": year, "주기": wlab, "응답주체": resp,
            "응답자ID": df[IDCOL[resp]].astype("Int64").astype(str),
            "학교ID": df["SCHID"].astype("Int64"),
            "indicator_id": cid, "원변수명": var, "값": s.values,
            "결측보정": ("미선택→0" if filled else "") + aligned})
        recs.append(sub)
micro = pd.concat(recs, ignore_index=True)
micro["지역규모코드"] = micro.apply(lambda r: REGION[1 if r["조사연도"]==2021 else 2].get(r["학교ID"]), axis=1)
micro["지역규모"] = micro["지역규모코드"].map(REGION_LBL).fillna("미상")

print("마이크로 long 행수:", len(micro))

# ── 3. 집계 ──────────────────────────────────────────────────
def agg(g, st, smin, smax):
    v = g.dropna()
    out = {"응답수": int(len(v)), "결측수": int(g.isna().sum())}
    if len(v)==0:
        return {**out, "평균":np.nan,"표준편차":np.nan,"중앙값":np.nan,"백분위25":np.nan,
                "백분위75":np.nan,"절사평균":np.nan,"환산100":np.nan,"상위2선지비율":np.nan,
                "긍정응답률":np.nan,"지출자수":np.nan,"지출참여율":np.nan,
                "지출자평균":np.nan,"지출자중앙값":np.nan}
    out.update(평균=round(float(v.mean()),4), 표준편차=round(float(v.std(ddof=1)),4) if len(v)>1 else np.nan,
               중앙값=float(v.median()), 백분위25=float(v.quantile(.25)), 백분위75=float(v.quantile(.75)))
    if st=="continuous":
        pay = v[v > 0]
        out["지출자수"] = int(len(pay))
        # 분모는 해당 주기 전체 응답자(결측 포함). 2주기 결측=미지출 가정 하에서 참여율이 되며,
        # 가정이 성립하지 않으면 하한값으로 해석한다.
        out["지출참여율"] = round(len(pay)/len(g)*100, 2) if len(g) else np.nan
        out["지출자평균"] = round(float(pay.mean()), 4) if len(pay) else np.nan
        out["지출자중앙값"] = float(pay.median()) if len(pay) else np.nan
    else:
        out["지출자수"] = out["지출참여율"] = out["지출자평균"] = out["지출자중앙값"] = np.nan
    if st=="continuous" and len(v)>=20:
        lo,hi = v.quantile(.01), v.quantile(.99)
        out["절사평균"] = round(float(v[(v>=lo)&(v<=hi)].mean()),4)
    else:
        out["절사평균"] = np.nan
    if st=="likert" and smin!="" and smax!="" and smax>smin:
        out["환산100"] = round((float(v.mean())-smin)/(smax-smin)*100, 2)
        top2 = [smax, smax-1]
        out["상위2선지비율"] = round(float(v.isin(top2).mean()*100), 2)
    else:
        out["환산100"], out["상위2선지비율"] = np.nan, np.nan
    return out

def build_facts(by_region: bool):
    rows = []
    keys = ["indicator_id","조사연도","주기"] + (["지역규모코드","지역규모"] if by_region else [])
    for k, g in micro.groupby(keys, dropna=False):
        cid = k[0]
        ix = dim.loc[dim.indicator_id==cid].iloc[0]
        if by_region and (pd.isna(k[3])): continue
        st = ix["통계유형"]
        smin = ix["척도최소"] if ix["척도최소"]!="" else ""
        smax = ix["척도최대"] if ix["척도최대"]!="" else ""
        rec = dict(zip(keys, k))
        rec.update(조사명="경기학교교육실태조사", 학교급="중학교")
        rec.update(대분류코드=ix["대분류코드"], 대분류=ix["대분류"], 소주제코드=ix["소주제코드"],
                   소주제=ix["소주제"], 응답주체=ix["응답주체"], 지표명=ix["지표명"],
                   통계유형=st, 시계열비교가능=ix["시계열비교가능"])
        rec.update(agg(g["값"], st, smin, smax))
        # 대표값: 지표 유형과 무관하게 그래프 y축에 바로 쓰는 단일 값
        if st=="likert":
            rec["대표값"], rec["대표값유형"] = rec.get("환산100"), "환산100점"
        elif st=="binary":
            rec["대표값유형"] = "긍정응답률(%)"
        elif st=="continuous":
            rec["대표값"] = rec.get("지출자중앙값") if pd.notna(rec.get("지출자중앙값")) else rec.get("중앙값")
            rec["대표값유형"] = "지출자 중앙값(만원)"
        else:
            rec["대표값"], rec["대표값유형"] = np.nan, ""
        if st=="binary" and ix["긍정코드"]!="":
            v = g["값"].dropna()
            rec["긍정응답률"] = round(float((v==ix["긍정코드"]).mean()*100),2) if len(v) else np.nan
            rec["대표값"] = rec["긍정응답률"]
        else:
            rec.setdefault("긍정응답률", np.nan)
            rec.setdefault("대표값", np.nan)
        rows.append(rec)
    return pd.DataFrame(rows)

_wave   = build_facts(False)
_region = build_facts(True)
_wave.insert(3, "지역규모코드", 0)
_wave.insert(4, "지역규모", "전체")
# 업체는 이 한 테이블만 물리면 된다.
#   연도별 추이   → 지역규모='전체' 필터
#   지역규모 비교 → 지역규모<>'전체' 필터
fact_wave = pd.concat([_wave, _region], ignore_index=True)
fact_region = _region   # 하위 호환(별도 파일로도 유지)

# 선지 분포
dist = (micro[micro["값"].notna()]
        .merge(dim[["indicator_id","통계유형","대분류코드","대분류","소주제코드","소주제","지표명"]],
               on="indicator_id"))
dist = dist[dist["통계유형"].isin(["likert","binary"])]
d = (dist.groupby(["indicator_id","대분류코드","대분류","소주제코드","소주제","응답주체","지표명","조사연도","주기","값"])
        .size().reset_index(name="응답수"))
tot = d.groupby(["indicator_id","조사연도"])["응답수"].transform("sum")
d["비율"] = (d["응답수"]/tot*100).round(2)
vl = pd.DataFrame(vlab_rows)
d = d.merge(vl.rename(columns={"코드":"값","라벨":"선지라벨"}), on=["indicator_id","조사연도","값"], how="left")
fact_dist = d.rename(columns={"값":"코드"})
fact_dist.insert(0,"학교급","중학교"); fact_dist.insert(0,"조사명","경기학교교육실태조사")

# 소주제 총점(리커트 환산100 평균, 역문항·비교불가 제외)
core = fact_wave[(fact_wave.통계유형=="likert") & (fact_wave.시계열비교가능.isin(["Y","조건부"]))]
core = core[~core.indicator_id.isin(REVERSE_ITEMS | SCALE_MISFIT)]
fact_topic = (core.groupby(["대분류코드","대분류","소주제코드","소주제","응답주체","조사연도","주기"])
    .agg(지표수=("indicator_id","nunique"), 평균환산100=("환산100","mean"),
         평균응답수=("응답수","mean")).reset_index())
fact_topic["평균환산100"] = fact_topic["평균환산100"].round(2)
fact_topic["평균응답수"] = fact_topic["평균응답수"].round(0)
fact_topic.insert(0,"학교급","중학교"); fact_topic.insert(0,"조사명","경기학교교육실태조사")

# 헤드라인 카드 (2021→2025 증감)
p = fact_topic.pivot_table(index=["대분류코드","대분류","소주제코드","소주제","응답주체"],
                           columns="조사연도", values="평균환산100").reset_index()
# 비교 시점은 WAVES의 마지막 두 개를 사용한다(주기 추가 시 자동으로 최신 두 시점).
_years = sorted(y for _, _, y, _ in WAVES)
PREV, CURR = _years[-2], _years[-1]
p = p.dropna(subset=[PREV, CURR])
p["직전시점"], p["최신시점"] = PREV, CURR
p["증감"] = (p[CURR]-p[PREV]).round(2)
p["증감률"] = (p["증감"]/p[PREV]*100).round(2)
p["변화방향"] = np.where(p["증감"]>0,"증가",np.where(p["증감"]<0,"감소","유지"))
p["절대증감순위"] = p["증감"].abs().rank(ascending=False, method="min").astype(int)
fact_headline = p.rename(columns={PREV:f"값_{PREV}", CURR:f"값_{CURR}"}).sort_values("절대증감순위")
fact_headline.insert(0,"학교급","중학교"); fact_headline.insert(0,"조사명","경기학교교육실태조사")

# ── 4. 품질 점검 ─────────────────────────────────────────────
qc = fact_wave.pivot_table(index="indicator_id", columns="조사연도",
                           values=["응답수","결측수","평균"]).reset_index()
qc.columns = ["indicator_id"] + [f"{a}_{b}" for a,b in qc.columns[1:]]
qc = dim[["indicator_id","대분류","소주제","응답주체","지표명","통계유형","시계열비교가능",
          "역문항","데이터이슈","응답기저","주의사항"]].merge(qc, on="indicator_id", how="left")
excl = dim[dim.시계열비교가능=="N"][["indicator_id","대분류","소주제","응답주체","지표명",
        "연계판정","통계유형","데이터이슈","주의사항"]].copy()
excl["제외사유"] = np.where(excl.데이터이슈!="", "2주기 코딩 이상(확인 필요)",
                      np.where(excl.연계판정.str.startswith("③"), "측정방식 상이 — 비교 불가",
                      np.where(excl.연계판정.str.startswith("④"), "한 주기에만 존재", "기타")))

# ── 5. 저장 ──────────────────────────────────────────────────
def w(df, path, **kw):
    df.to_csv(os.path.join(OUT,path), index=False, encoding="utf-8-sig", **kw)
    print(f"  {path}: {len(df):,}행")

w(dim, "dim_indicator.csv")
w(pd.DataFrame(CATEGORIES, columns=["대분류코드","대분류","출처","주요 응답주체"]), "dim_category.csv")
w(dim[["대분류코드","대분류","소주제코드","소주제","응답주체"]].drop_duplicates()
     .sort_values(["대분류코드","소주제코드"]), "dim_topic.csv")
w(pd.DataFrame([{"조사연도":c,"주기코드":a,"주기번호":b,"주기명":d,
                 "조사명":"경기학교교육실태조사","학교급":"중학교"} for a,b,c,d in WAVES]),
  "dim_wave.csv")
w(pd.DataFrame([{"지역규모코드":k,"지역규모":v} for k,v in REGION_LBL.items()]), "dim_region.csv")
w(pd.DataFrame(vlab_rows).drop_duplicates(), "dim_value_label.csv")

# 응답주체 간 비교(다계열 그래프) 성립 여부 — 변수매칭표 5번 시트
_cr = [[N(str(c)) if c is not None else "" for c in r]
       for r in wb["5_응답주체간비교"].iter_rows(values_only=True)][2:]
_cr = [r for r in _cr if r[0]]
cross = pd.DataFrame([dict(
    등급=r[0], 성립여부=("성립" if r[1]=="유지" else "불가"), 사유=(""if r[1]=="유지" else r[1]),
    비교조합=r[2], 주제=r[3],
    학생_1주기=r[4], 학생_2주기=r[5], 학생판정=r[6],
    학부모_1주기=r[7], 학부모_2주기=r[8], 학부모판정=r[9],
    교사_1주기=r[10], 교사_2주기=r[11], 교사판정=r[12],
    척도=r[13], 주의사항=r[14]) for r in _cr])
w(cross, "dim_cross_respondent.csv")
sch = pd.concat([
  pd.DataFrame({"조사연도":2021,"학교ID":db1.SCHID.astype(int),"지역규모코드":db1.YM1_DB0_3,
                "설립구분코드":db1.YM1_DB0_1,"남녀공학코드":db1.YM1_DB0_4}),
  pd.DataFrame({"조사연도":2025,"학교ID":db2.SCHID.astype(int),"지역규모코드":db2.YM2_DB0_3,
                "설립구분코드":db2.YM2_DB0_1,"남녀공학코드":db2.YM2_DB0_4})])
sch["지역규모"] = sch["지역규모코드"].map(REGION_LBL)
sch["설립구분"] = sch["설립구분코드"].map({1.0:"공립",2.0:"사립"})
sch["남녀공학"] = sch["남녀공학코드"].map({1.0:"남학교",2.0:"여학교",3.0:"남여공학"})
# 학교DB에는 있으나 실제 응답이 수집되지 않은 학교가 있다. 응답주체별로 표기한다.
for resp, col in [("학생","학생응답"),("학부모","학부모응답"),("교사","교사응답")]:
    have = micro[micro.응답주체==resp].groupby("조사연도")["학교ID"].apply(set).to_dict()
    sch[col] = [("Y" if k in have.get(y,set()) else "N")
                for y,k in zip(sch["조사연도"], sch["학교ID"])]
w(sch, "dim_school.csv")

ORDER = ["조사명","학교급","조사연도","주기","지역규모코드","지역규모",
         "대분류코드","대분류","소주제코드","소주제","응답주체","indicator_id","지표명",
         "통계유형","시계열비교가능","대표값","대표값유형",
         "응답수","결측수","평균","표준편차","중앙값","백분위25","백분위75",
         "환산100","상위2선지비율","긍정응답률",
         "절사평균","지출자수","지출참여율","지출자평균","지출자중앙값"]
def reorder(df):
    cols=[c for c in ORDER if c in df.columns]+[c for c in df.columns if c not in ORDER]
    return df[cols]
fact_wave = reorder(fact_wave)
w(fact_wave.sort_values(["대분류코드","소주제코드","indicator_id","조사연도","지역규모코드"]),
  "fact_indicator_wave.csv")
w(reorder(fact_region).sort_values(["대분류코드","소주제코드","indicator_id","조사연도","지역규모코드"]),
  "fact_indicator_wave_region.csv")
w(fact_dist.sort_values(["indicator_id","조사연도","코드"]), "fact_indicator_dist.csv")
w(fact_topic.sort_values(["대분류코드","소주제코드","조사연도"]), "fact_topic_wave.csv")
w(fact_headline, "fact_headline.csv")

# 조사명·학교급·주기·원변수명·지역규모는 dim_wave / dim_indicator / dim_region에서 조회 가능하므로
# 마이크로데이터에서는 제외한다(파일 크기 58% 감소).
# 압축하지 않고 평문 CSV로 두며, 주기별로 나누어 모든 파일이 엑셀 행 한도 내에 들어가게 한다.
MICRO_COLS = ["조사연도","응답주체","응답자ID","학교ID","indicator_id","값","결측보정","지역규모코드"]
for resp, fn in [("학생","student"),("학부모","parent"),("교사","teacher")]:
    for _, _, year, _ in WAVES:
        sub = micro[(micro.응답주체==resp) & (micro.조사연도==year)][MICRO_COLS]
        if sub.empty: continue
        name = f"micro_{fn}_{year}.csv"
        sub.to_csv(os.path.join(OUT, name), index=False, encoding="utf-8-sig")
        print(f"  {name}: {len(sub):,}행")



import tabledef, guidedoc
_xl, _n = tabledef.build(OUT)
guidedoc.build(OUT, dim, micro, fact_wave, fact_topic, fact_headline, cross, WAVES)
print("  00_전달안내.md")
print(f"  01_정의/테이블정의서.xlsx  " + " / ".join(f"{k}: {v}행" for k, v in _n.items()))

print("\n완료:", OUT)
