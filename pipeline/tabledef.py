# -*- coding: utf-8 -*-
"""테이블 정의서 생성 — 산출된 CSV를 직접 읽어 스키마를 기술한다.
build.py 마지막 단계에서 호출되므로 데이터와 항상 일치한다."""
import os, glob, gzip
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# 컬럼명은 한글을 정식 명칭으로 사용한다(2026-09-07 결정).
# 아래 영문명은 BI 도구·DB가 한글 컬럼을 지원하지 않을 경우를 대비한 예비 매핑이며,
# 전환이 필요하면 이 딕셔너리로 CSV 헤더를 치환한 뒤 재생성하면 된다.
EN = {
 "조사명":"survey_name","학교급":"school_level","조사연도":"survey_year","주기":"wave_name",
 "주기코드":"wave_code","주기번호":"wave_no","주기명":"wave_label",
 "지역규모코드":"region_code","지역규모":"region_name",
 "대분류코드":"category_code","대분류":"category_name",
 "소주제코드":"topic_code","소주제":"topic_name","응답주체":"respondent",
 "indicator_id":"indicator_id","지표명":"indicator_name","문항명":"question_name",
 "하위문항":"item_name","원영역":"source_area","출처":"source",
 "var_2021":"var_w1","var_2025":"var_w2","연계판정":"link_grade","통계유형":"stat_type",
 "선지수":"scale_points","척도최소":"scale_min","척도최대":"scale_max",
 "척도최소라벨":"scale_min_label","척도최대라벨":"scale_max_label",
 "긍정코드":"positive_code","역문항":"reverse_item","시계열비교가능":"comparable",
 "주체간비교":"cross_respondent","확인필요":"needs_check","데이터이슈":"data_issue",
 "응답기저":"response_base","주의사항":"caution",
 "대표값":"headline_value","대표값유형":"headline_value_type",
 "응답수":"n_valid","결측수":"n_missing","평균":"mean","표준편차":"sd","중앙값":"median",
 "백분위25":"p25","백분위75":"p75","절사평균":"trimmed_mean",
 "환산100":"scaled_100","상위2선지비율":"top2_pct","긍정응답률":"positive_pct",
 "지출자수":"n_payers","지출참여율":"payer_pct","지출자평균":"payer_mean",
 "지출자중앙값":"payer_median",
 "성립여부":"is_valid","사유":"invalid_reason","비교조합":"pair","주제":"topic",
 "학생응답":"has_student","학부모응답":"has_parent","교사응답":"has_teacher",
 "직전시점":"prev_year","최신시점":"curr_year","등급":"grade","척도":"scale",
 "코드":"code","라벨":"label","선지라벨":"option_label","비율":"pct",
 "학교ID":"school_id","설립구분코드":"found_code","설립구분":"found_name",
 "남녀공학코드":"coed_code","남녀공학":"coed_name",
 "값_2021":"value_w1","값_2025":"value_w2","증감":"delta","증감률":"delta_pct",
 "변화방향":"direction","절대증감순위":"abs_delta_rank",
 "지표수":"n_indicators","평균환산100":"mean_scaled_100","평균응답수":"mean_n_valid",
 "응답자ID":"respondent_id","원변수명":"source_var","값":"value","결측보정":"missing_fix",
 "주요 응답주체":"main_respondents",
}

DESC = {
 "조사명":"고정값 '경기학교교육실태조사'. 향후 타 조사 확장 대비",
 "학교급":"고정값 '중학교'. 초·고 확장 시 값이 늘어남 (화면 ① 전역 필터)",
 "조사연도":"2021(1주기) 또는 2025(2주기)",
 "주기":"'1주기' 또는 '2주기'",
 "지역규모코드":"0=전체, 1=중소도시, 2=읍면지역",
 "지역규모":"'전체' / '중소도시' / '읍면지역'. ⑤-① 추이는 '전체', ⑤-② 비교는 '전체' 제외",
 "대분류코드":"C01~C12. 화면 ② 대분류 탭",
 "대분류":"대분류 표시명",
 "소주제코드":"S로 시작. 화면 ③ 소주제 탭",
 "소주제":"소주제 표시명",
 "응답주체":"'학생' / '학부모' / '교사'. 다계열 그래프의 계열 구분",
 "indicator_id":"지표 고유 식별자. 주기가 바뀌어도 불변 (전 테이블 조인 키)",
 "지표명":"화면 표시용 지표명 (문항명 - 하위문항)",
 "대표값":"★ 그래프 y축에 바로 사용. 지표 유형과 무관하게 채워짐",
 "대표값유형":"★ 대표값의 단위·의미. 축 이름 표기에 사용",
 "응답수":"유효 응답자 수. 그래프에 병기 필요",
 "결측수":"결측 응답자 수",
 "평균":"원 척도 기준 평균",
 "표준편차":"표본표준편차",
 "중앙값":"중앙값",
 "환산100":"(평균-척도최소)/(척도최대-척도최소)×100. 선택형 지표만",
 "상위2선지비율":"최상위 2개 선지 응답 비율(%). 선택형 지표만",
 "긍정응답률":"긍정코드 응답 비율(%). 이분형 지표만",
 "지출자중앙값":"0 초과 응답자의 중앙값. 금액형 지표만",
 "시계열비교가능":"Y=그대로 비교 / 조건부=주의사항 확인 / N=시계열 제외",
 "통계유형":"likert(선택형) / binary(이분형) / continuous(금액형) / excluded",
 "역문항":"Y=역채점 필요 / 척도 부적합 / N",
 "var_2021":"1주기 SPSS 원변수명",
 "var_2025":"2주기 SPSS 원변수명",
 "코드":"응답 선지 코드값",
 "라벨":"해당 코드의 의미. 주기별로 표현이 다를 수 있음",
 "값":"응답값 (코드 또는 실수)",
 "결측보정":"적용된 보정 내용. 공란이면 원값 그대로",
 "성립여부":"'성립'이면 다계열 그래프 구성 가능. '불가'는 한쪽 주체 문항이 미선택되어 단일 계열로 표시",
 "사유":"성립 불가 사유. 어느 주체가 미선택인지 표기",
 "학생응답":"해당 주기에 이 학교에서 학생 응답이 수집되었는지 (Y/N)",
 "학부모응답":"해당 주기에 이 학교에서 학부모 응답이 수집되었는지 (Y/N)",
 "교사응답":"해당 주기에 이 학교에서 교사 응답이 수집되었는지 (Y/N)",
 "직전시점":"증감 계산의 기준이 된 이전 조사연도",
 "최신시점":"증감 계산의 기준이 된 최근 조사연도",
 "응답자ID":"학생·학부모는 STUID(학부모는 자녀의 학생 ID), 교사는 TID. 집계에는 사용하지 않는다. "
          "동일 주기의 학생·학부모는 같은 값으로 연결되나, 주기가 다르면 서로 다른 사람이다",
 "학교ID":"SCHID. 주기 간 동일 학교를 가리킴",
}

TABLES = [
 ("dim_category.csv","dim_category","차원","대분류 목록",
  "화면 ② 대분류 탭 구성"),
 ("dim_topic.csv","dim_topic","차원","대분류–소주제 트리",
  "화면 ③ 소주제 탭 구성"),
 ("dim_indicator.csv","dim_indicator","차원","지표 마스터",
  "지표 정의·주기별 원변수·척도·주의사항. 전 테이블의 기준"),
 ("dim_value_label.csv","dim_value_label","차원","선지 코드–라벨",
  "화면 ⑤-③ 분포 그래프의 범례. 주기별로 라벨이 다를 수 있음"),
 ("dim_wave.csv","dim_wave","차원","조사 주기",
  "주기 추가 시 이 표에 행을 추가"),
 ("dim_region.csv","dim_region","차원","지역규모 코드",""),
 ("dim_school.csv","dim_school","차원","학교 속성",
  "지역규모 결합 근거. 응답 수집 여부를 응답주체별로 표기"),
 ("dim_cross_respondent.csv","dim_cross_respondent","차원","응답주체 간 비교 성립 여부",
  "화면 ⑤ 다계열 그래프. 27개 조합 중 12개만 성립하므로 반드시 확인"),
 ("fact_indicator_wave.csv","fact_indicator_wave","사실","지표별 집계 (전체·지역 통합)",
  "★ 주력 테이블. 화면 ⑤-① 추이와 ⑤-② 지역 비교를 모두 커버"),
 ("fact_topic_wave.csv","fact_topic_wave","사실","소주제별 집계",
  "화면 ④ 헤드라인 카드의 소주제 값"),
 ("fact_headline.csv","fact_headline","사실","소주제 증감 순위",
  "화면 ④ 헤드라인 카드"),
 ("fact_indicator_dist.csv","fact_indicator_dist","사실","선지별 응답 분포",
  "화면 ⑤-③ 개별 문항 2시점 비교"),
 ("fact_indicator_wave_region.csv","fact_indicator_wave_region","사실",
  "지표별 지역규모 집계 (분리본)","fact_indicator_wave의 지역 부분. 선택 사용"),
 ("micro_student_2021.csv","micro_student_2021","원자료","학생 개인 단위 long (1주기)",""),
 ("micro_student_2025.csv","micro_student_2025","원자료","학생 개인 단위 long (2주기)",""),
 ("micro_parent_2021.csv","micro_parent_2021","원자료","학부모 개인 단위 long (1주기)",""),
 ("micro_parent_2025.csv","micro_parent_2025","원자료","학부모 개인 단위 long (2주기)",""),
 ("micro_teacher_2021.csv","micro_teacher_2021","원자료","교사 개인 단위 long (1주기)",""),
 ("micro_teacher_2025.csv","micro_teacher_2025","원자료","교사 개인 단위 long (2주기)",
  "집계 재현 및 연도 갱신용. 응답주체·주기별로 6개 파일"),
]

JOINS = [
 ("fact_indicator_wave","dim_indicator","indicator_id","N:1","지표 정의·주의사항 조회"),
 ("fact_indicator_wave","dim_category","대분류코드","N:1","대분류 탭"),
 ("fact_indicator_wave","dim_topic","대분류코드 + 소주제코드","N:1","소주제 탭"),
 ("fact_indicator_wave","dim_region","지역규모코드","N:1","지역 필터"),
 ("fact_indicator_wave","dim_wave","조사연도","N:1","주기 정보"),
 ("fact_indicator_dist","dim_value_label","indicator_id + 조사연도 + 코드","N:1","선지 라벨"),
 ("fact_topic_wave","dim_topic","대분류코드 + 소주제코드","N:1","소주제 정보"),
 ("fact_headline","dim_topic","대분류코드 + 소주제코드","N:1","카드 클릭 시 이동 대상"),
 ("micro_*","dim_indicator","indicator_id","N:1","지표 정의"),
 ("micro_*","dim_school","조사연도 + 학교ID","N:1","학교 속성"),
]

def dtype_of(s):
    if pd.api.types.is_integer_dtype(s) or (pd.api.types.is_float_dtype(s)
        and s.dropna().mod(1).eq(0).all() and s.notna().any()): return "정수"
    if pd.api.types.is_float_dtype(s): return "실수"
    return "문자"

def build(out_dir):
    tl, cl = [], []
    for path, phys, kind, name, note in TABLES:
        full = os.path.join(out_dir, path)
        if not os.path.exists(full): continue
        if path.endswith(".gz"):
            df = pd.read_csv(full, nrows=3000)
            with gzip.open(full, "rt", encoding="utf-8-sig") as fh:
                n = sum(1 for _ in fh) - 1
        else:
            df = pd.read_csv(full)
            n = len(df)
        tl.append(dict(구분=kind, 테이블명=phys, 논리명=name, 파일=path,
                       행수=n, 열수=len(df.columns), 용도=note))
        for i, c in enumerate(df.columns, 1):
            cl.append(dict(테이블명=phys, 순번=i, 컬럼명=c, **{"영문명(예비)": EN.get(c, "")},
                           자료형=dtype_of(df[c]),
                           널허용="N" if df[c].notna().all() else "Y",
                           설명=DESC.get(c, ""),
                           예시=str(df[c].dropna().iloc[0])[:38] if df[c].notna().any() else ""))
    sheets = {
        "1. 테이블 목록": pd.DataFrame(tl),
        "2. 컬럼 정의": pd.DataFrame(cl),
        "3. 테이블 관계": pd.DataFrame(JOINS, columns=["기준 테이블","참조 테이블","조인 키","관계","용도"]),
        "4. 명명·형식 규칙": pd.DataFrame([
          dict(항목="컬럼명", 규칙="한글 사용",
               내용="CSV 헤더는 한글을 정식 명칭으로 한다. '2. 컬럼 정의' 시트의 '영문명(예비)'은 "
                    "BI 도구·DB가 한글 컬럼을 지원하지 않을 경우를 위한 대체안이며 현재 파일에는 쓰이지 않는다."),
          dict(항목="파일 형식", 규칙="CSV (UTF-8 BOM)",
               내용="Excel에서 바로 열리도록 BOM을 포함한다. 구분자는 쉼표, 인용부호는 큰따옴표."),
          dict(항목="테이블명", 규칙="영문 스네이크",
               내용="dim_* 은 차원(변하지 않는 정의), fact_* 은 사실(측정값), micro_* 은 개인 단위 원자료."),
          dict(항목="결측 표현", 규칙="빈칸",
               내용="해당 지표 유형에 없는 통계량은 빈칸이다(예: 선택형 지표의 '지출자평균'). "
                    "'대표값' 컬럼만 사용하면 지표 유형별 분기가 불필요하다."),
          dict(항목="수치 자리수", 규칙="원값 보존",
               내용="평균 소수 4자리, 환산100·비율 2자리로 저장한다. 화면 표시용 반올림은 업체 측에서 처리한다."),
          dict(항목="코드값", 규칙="코드+명칭 병기",
               내용="지역규모·대분류·소주제 등은 코드 컬럼과 명칭 컬럼을 함께 제공한다. "
                    "정렬·필터는 코드, 표시는 명칭을 사용하면 된다."),
          dict(항목="갱신 방식", 규칙="전체 교체",
               내용="주기가 추가되면 전 파일을 재생성하여 교체한다. 증분 갱신은 지원하지 않는다."),
        ]),
    }
    xl = os.path.join(out_dir, "01_테이블정의서.xlsx")
    with pd.ExcelWriter(xl, engine="openpyxl") as xw:
        for k, v in sheets.items(): v.to_excel(xw, sheet_name=k, index=False)
        wb = xw.book
        HDR = PatternFill("solid", fgColor="DDE4EE")
        line = Side(style="thin", color="BFC8D6"); hair = Side(style="hair", color="D8DEE8")
        W = {"구분":7,"테이블명":26,"논리명":26,"파일":34,"행수":9,"열수":5,"용도":46,
             "순번":5,"컬럼명":18,"영문명(예비)":20,"자료형":7,"널허용":7,"설명":58,"예시":30,
             "기준 테이블":26,"참조 테이블":22,"조인 키":30,"관계":6}
        for k, v in sheets.items():
            ws = wb[k]; ws.freeze_panes = "A2"; ws.sheet_view.showGridLines = False
            for c in ws[1]:
                c.font = Font(bold=True, size=10, color="1F2A3C"); c.fill = HDR
                c.alignment = Alignment(vertical="center", wrap_text=True)
                c.border = Border(top=line, bottom=line)
            ws.row_dimensions[1].height = 24
            for i, col in enumerate(v.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = W.get(col, 20)
            last = ws.max_row
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(vertical="top", wrap_text=True); c.font = Font(size=10)
                    c.border = Border(bottom=(line if row[0].row == last else hair))
    return xl, {k: len(v) for k, v in sheets.items()}
